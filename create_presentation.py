import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = RGBColor(8, 14, 24)
    CARD_BG = RGBColor(17, 28, 48)
    CARD_BORDER = RGBColor(30, 48, 77)
    
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_LIGHT = RGBColor(203, 213, 225)
    TEXT_MUTED = RGBColor(148, 163, 184)
    
    EMERALD = RGBColor(16, 185, 129)
    TEAL = RGBColor(20, 184, 166)
    AMBER = RGBColor(245, 158, 11)
    ROSE = RGBColor(239, 68, 68)
    INDIGO = RGBColor(99, 102, 241)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="AAROGYAVANI • HEALTHCARE AI", page_num=None):
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = category_text.upper()
        p_badge.font.size = Pt(10)
        p_badge.font.bold = True
        p_badge.font.color.rgb = EMERALD
        p_badge.font.name = "Segoe UI"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Segoe UI"

        if page_num:
            num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1), Inches(0.4))
            p_num = num_box.text_frame.paragraphs[0]
            p_num.text = f"{page_num} / 12"
            p_num.font.size = Pt(11)
            p_num.font.bold = True
            p_num.font.color.rgb = TEXT_MUTED
            p_num.font.name = "Segoe UI"
            p_num.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # SLIDE 1
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    orb = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(1.5), Inches(4.5), Inches(4.5))
    orb.fill.solid()
    orb.fill.fore_color.rgb = RGBColor(16, 185, 129)
    orb.line.fill.background()

    create_card(s1, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5), RGBColor(12, 21, 37), EMERALD)
    tb1 = s1.shapes.add_textbox(Inches(1.3), Inches(1.4), Inches(10.5), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "BIT N BUILD '26 HACKATHON  •  HEALTHTECH TRACK"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p = tf1.add_paragraph()
    p.text = "आरोग्यवाणी (AarogyaVani)"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Voice-First AI Clinical Triage & Referral Assistant for Rural Healthcare"
    p.font.size = Pt(20)
    p.font.color.rgb = TEAL
    p.font.bold = True
    p.space_after = Pt(24)

    p = tf1.add_paragraph()
    p.text = "Bridging the rural healthcare divide for 850M+ citizens and 1M+ ASHA workers.\nSpoken symptoms in native Indian languages  ➔  Deterministic danger safety  ➔  Color-coded triage + Voice guidance."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(36)

    p = tf1.add_paragraph()
    p.text = "Developed by Team LocalHost Boys • GDG HBTU  |  React 19 • Gemini 2.5 Flash • Sarvam AI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_MUTED

    # SLIDE 2
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem: The Rural Indian Healthcare Trilemma", "GROUND REALITY & CONTEXT", 2)
    card_w = Inches(3.64)
    card_h = Inches(5.1)
    top_pos = Inches(1.6)

    c1 = create_card(s2, Inches(0.8), top_pos, card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👨‍⚕️ 1. Extreme Scarcity"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ROSE
    p.space_after = Pt(12)
    bullets1 = [
        ("Doctor Ratio 1:25,000+", "Rural regions face severe doctor shortages, 25x worse than WHO's 1:1,000 baseline."),
        ("40 km Average Travel", "Villagers travel hours on foot or tractor to reach the nearest Community Health Centre (CHC)."),
        ("Late Clinical Arrival", "Patients arrive when self-limiting fevers become critical sepsis or heart failure.")
    ]
    for b_title, b_desc in bullets1:
        p = tf.add_paragraph()
        p.text = f"• {b_title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {b_desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    c2 = create_card(s2, Inches(4.84), top_pos, card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(5.04), top_pos + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🗣️ 2. Language & Literacy"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_after = Pt(12)
    bullets2 = [
        ("Keyboards Fail", "Typing symptoms in English or complex apps is impossible for low-literacy rural patients."),
        ("Hundreds of Dialects", "Rural speakers use Bhojpuri, Maithili, Awadhi, or code-mixed vernacular idioms."),
        ("Unreadable Prescriptions", "Old handwritten doctor slips are lost or unreadable, leading to repeated misdiagnosis.")
    ]
    for b_title, b_desc in bullets2:
        p = tf.add_paragraph()
        p.text = f"• {b_title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {b_desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    c3 = create_card(s2, Inches(8.88), top_pos, card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(9.08), top_pos + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👩‍⚕️ 3. ASHA Workload"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(12)
    bullets3 = [
        ("1 Million Workers", "ASHA workers are the frontline backbone but carry massive physical paper registers."),
        ("No Diagnostic Tools", "ASHAs have basic training but no specialist triage decision support at the doorstep."),
        ("Referral Bottlenecks", "PHC doctors receive uncertified verbal complaints without vitals or structured history.")
    ]
    for b_title, b_desc in bullets3:
        p = tf.add_paragraph()
        p.text = f"• {b_title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {b_desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    # SLIDE 3
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "The Solution: What is AarogyaVani?", "INNOVATION & VALUE PROPOSITION", 3)
    w_half = Inches(5.66)
    create_card(s3, Inches(0.8), Inches(1.6), w_half, Inches(5.1), RGBColor(14, 25, 45), EMERALD)
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(1.8), w_half - Inches(0.6), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "💡 For Non-Technical & Healthcare Judges"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(14)
    non_tech_points = [
        ("Virtual Clinical Copilot", "A compassionate, pocket-sized assistant for ASHA workers and villagers."),
        ("Speak Naturally", "Just press one big button and speak symptoms in your native tongue. No typing needed."),
        ("Understands Any Accent", "Processes Hindi, Bengali, Telugu, Marathi, and regional rural dialects."),
        ("Instant Danger Protection", "Detects heart attack, stroke, or severe bleeding in less than a second."),
        ("Voice-Out Guidance", "Calmly speaks out home care tips and instructions so illiterate patients understand.")
    ]
    for tit, dsc in non_tech_points:
        p = tf.add_paragraph()
        p.text = f"✔ {tit}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"   {dsc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    create_card(s3, Inches(6.86), Inches(1.6), w_half, Inches(5.1), RGBColor(14, 25, 45), TEAL)
    tb = s3.shapes.add_textbox(Inches(7.16), Inches(1.8), w_half - Inches(0.6), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ For Technical Evaluators & Architects"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(14)
    tech_points = [
        ("Multi-Model Pipeline", "Combines Sarvam Saaras v4 (ASR), Gemini 2.5 Flash (Triage), and Bulbul v3 (TTS)."),
        ("Deterministic Safety Gates", "Hardcoded regex rules bypass LLM to deliver 100% zero-hallucination RED alerts."),
        ("Multimodal Vision OCR", "Gemini Vision deciphers handwritten prescriptions into clean JSON in <2 seconds."),
        ("Resilient Fallback Design", "If Sarvam ASR fails, auto-falls back to Gemini Multimodal Audio without breaking."),
        ("Zero PII Retention", "Stateless serverless execution complies with healthcare privacy and ABDM standards.")
    ]
    for tit, dsc in tech_points:
        p = tf.add_paragraph()
        p.text = f"⚙ {tit}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"   {dsc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    # SLIDE 4
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Dual-Persona Architecture: Citizen Simplicity + ASHA Rigor", "USER EXPERIENCE & TARGET PERSONAS", 4)
    create_card(s4, Inches(0.8), Inches(1.6), w_half, Inches(5.1), RGBColor(14, 24, 40), TEAL)
    tb = s4.shapes.add_textbox(Inches(1.1), Inches(1.8), w_half - Inches(0.6), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🧑‍🌾 Citizen / Patient Mode (मरीज़)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(10)
    cit_features = [
        ("Primary Goal", "Immediate self-triage for rural villagers at home."),
        ("Zero Typing Required", "Big central microphone button + interactive tap-based Anatomical Body Map."),
        ("Dialect Friendly", "Patient describes illness in colloquial Hindi, Bengali, Telugu, Marathi."),
        ("Instant 108 Emergency Dial", "One-tap direct dialer button connecting directly to state ambulance service."),
        ("Audio-First Output", "Plays spoken voice guidance automatically — no need to read medical words."),
        ("Reassurance & Privacy", "Zero registration friction, zero personal records stored on public devices.")
    ]
    for tit, dsc in cit_features:
        p = tf.add_paragraph()
        p.text = f"• {tit}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {dsc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(5)

    create_card(s4, Inches(6.86), Inches(1.6), w_half, Inches(5.1), RGBColor(14, 24, 40), EMERALD)
    tb = s4.shapes.add_textbox(Inches(7.16), Inches(1.8), w_half - Inches(0.6), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👩‍⚕️ ASHA Field Worker Mode (आशा)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(10)
    asha_features = [
        ("Primary Goal", "Frontline doorstep clinical screening and certified PHC referral."),
        ("Beneficiary Registration", "Quick capture of Patient Name, Age, and Gender for clinical record."),
        ("Vital Signs Gauges", "Optional input for Temperature (°F), Pulse (BPM), BP, and SpO2 (%)."),
        ("Prescription Scanner", "Laser-guided camera scan to digitize previous clinic prescriptions and reports."),
        ("Official NHM Referral Slip", "Formats formatted slip with clinical reasons and recommended next steps."),
        ("ASHA Sign-Off Badge", "Interactive 'Verified by ASHA ✓' verification that PHC doctors respect.")
    ]
    for tit, dsc in asha_features:
        p = tf.add_paragraph()
        p.text = f"• {tit}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {dsc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(5)

    # SLIDE 5
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "End-to-End System Flow: From Voice Note to Referral Slip", "THE CLINICAL PIPELINE", 5)
    step_w = Inches(2.7)
    step_h = Inches(5.1)
    top_step = Inches(1.6)

    steps = [
        ("1. INTAKE", "🎤 Multimodal Intake", [
            ("Voice Ingestion", "Patient/ASHA records spoken complaint (or uploads audio)."),
            ("Body Map Touch", "Selects affected regions (Chest, Head, Stomach, Arms)."),
            ("Document Scan", "Optional doctor prescription photo uploaded for OCR."),
            ("Clinical Vitals", "Temperature, Pulse, Blood Pressure, and Oxygen entered.")
        ], TEAL),
        ("2. SAFETY GATE", "🚨 Deterministic Gate", [
            ("Instant Keyword Match", "Regex checks 15+ acute clinical red flags in <100ms."),
            ("Zero LLM Latency", "Chest pain, seizures, or hemorrhage bypass all AI models."),
            ("Direct Emergency Path", "Immediate RED classification with 108 ambulance hotline."),
            ("Guaranteed Safety", "100% eliminates LLM hallucinations on lethal conditions.")
        ], ROSE),
        ("3. AI REASONING", "🧠 Clinical Stratification", [
            ("Gemini 2.5 Flash", "Receives structured prompt with symptoms, map, vitals & OCR."),
            ("WHO IMNCI Protocols", "Stratifies into YELLOW (Consult PHC) or GREEN (Routine)."),
            ("Explainable Reasons", "Generates bulleted rationale and actionable home-care steps."),
            ("Dual Language Output", "Returns structured text in English and Devanagari Hindi.")
        ], AMBER),
        ("4. DELIVERY", "🔊 Vernacular Action", [
            ("Sarvam Bulbul TTS", "Synthesizes native Indian voice advice in ~1.2 seconds."),
            ("Voice Player", "Patient taps Play to listen to calm, reassuring audio guidance."),
            ("NHM Referral Slip", "Structured medical summary ready to print or WhatsApp."),
            ("ASHA Verification", "Frontline certification logged for Sub-Centre / PHC doctor.")
        ], EMERALD),
    ]

    for idx, (badge, title, points, color) in enumerate(steps):
        left_pos = Inches(0.8) + idx * Inches(2.95)
        create_card(s5, left_pos, top_step, step_w, step_h, CARD_BG, color)
        tb = s5.shapes.add_textbox(left_pos + Inches(0.15), top_step + Inches(0.15), step_w - Inches(0.3), step_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = badge
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = color
        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)
        for p_tit, p_desc in points:
            p = tf.add_paragraph()
            p.text = f"• {p_tit}"
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE
            p = tf.add_paragraph()
            p.text = f"  {p_desc}"
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(6)

    # SLIDE 6
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Architectural Pipeline: Resilient Multi-Model Orchestration", "TECHNICAL ARCHITECTURE & APIS", 6)

    layers = [
        ("Layer 1: Edge Client UI", "React 19 + TypeScript + Vite + Tailwind CSS", 
         "• Runs responsive PWA across low-end Android mobile devices and desktop browsers.\n• Local audio recording via Web Audio API with automatic WAV/WebM containerization.\n• Zero-dependency i18n localization engine supporting Hindi, Bengali, Telugu, Marathi, English.", TEAL),
        ("Layer 2: Serverless Gateway", "Netlify Functions v2 (TypeScript Serverless)", 
         "• Stateless endpoints (/voice, /scan, /check) with CORS, rate-limiting, and error shielding.\n• Zero disk persistence: Audio buffers and OCR images processed ephemerally in RAM.\n• Environment credential isolation protecting Sarvam and Gemini API keys.", INDIGO),
        ("Layer 3: Speech AI Engine", "Sarvam Saaras v4 ASR ⇄ Gemini Multimodal Audio", 
         "• Primary: Sarvam Saaras v4 transcribes Hindi and Indic dialects with high acoustic accuracy.\n• Resilient Fallback: If Sarvam fails, Gemini Multimodal Audio transcribes code-mixed dialects.\n• Normalizes vernacular voice inputs into structured Hindi & English text representations.", EMERALD),
        ("Layer 4: Vision & Reasoning", "Gemini Vision OCR ⇄ Gemini 2.5 Flash Clinical Triage", 
         "• High-speed multimodal OCR deciphers doctor handwriting, dosages, and patient notes.\n• Strict JSON schema output: Enforces {level, reasons, guidance, guidanceHindi, nextSteps}.\n• Multi-model retry ladder: Automatically cycles across Flash-Lite and Flash if rate-limited.", AMBER),
        ("Layer 5: Voice Synthesis", "Sarvam Bulbul v3 TTS (Indic Audio Delivery)", 
         "• Generates natural Hindi speech ('priya' voice) from the AI's guidance text at pace 0.9.\n• Encodes audio stream into base64 payload delivered directly to client VoicePlayer component.\n• Average end-to-end synthesis roundtrip under 1.5 seconds.", ROSE)
    ]
    for idx, (l_title, l_sub, l_desc, l_color) in enumerate(layers):
        l_top = Inches(1.5) + idx * Inches(1.05)
        create_card(s6, Inches(0.8), l_top, Inches(11.733), Inches(0.95), CARD_BG, l_color)
        tb = s6.shapes.add_textbox(Inches(1.0), l_top + Inches(0.08), Inches(11.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{l_title}  |  {l_sub}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = l_color
        p = tf.add_paragraph()
        p.text = l_desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT

    # SLIDE 7
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Deterministic Safety Gates: <100ms Zero-Hallucination Path", "CLINICAL SAFETY & ETHICS", 7)
    create_card(s7, Inches(0.8), Inches(1.6), Inches(5.66), Inches(5.1), RGBColor(18, 22, 38), ROSE)
    tb = s7.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🛡️ Why Never Gamble with an LLM in Emergencies"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ROSE
    p.space_after = Pt(12)
    phil_pts = [
        ("Probabilistic LLMs are Risky", "Even a 1% hallucination rate on an acute myocardial infarction or neonatal seizure can prove fatal in rural areas."),
        ("Speed Matters in Sepsis & Trauma", "Waiting 3-5 seconds for an LLM token generation is unacceptable when a patient is bleeding out or unconscious."),
        ("Deterministic Regex Fast-Path", "AarogyaVani catches acute danger keywords in <100ms and immediately halts further processing."),
        ("Direct Ambulance Trigger", "Instantly locks UI to RED status, highlights 108 Hotline, and plays urgent spoken warning.")
    ]
    for tit, dsc in phil_pts:
        p = tf.add_paragraph()
        p.text = f"✔ {tit}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {dsc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    create_card(s7, Inches(6.86), Inches(1.6), Inches(5.66), Inches(5.1), CARD_BG, CARD_BORDER)
    tb = s7.shapes.add_textbox(Inches(7.16), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📋 Acute Danger Sign Catalog (Bypasses LLM)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(10)
    rules_table = [
        ("Chest Pain / Angina", "seene mein dard, chaati mein dard", "🔴 Immediate hospital transfer"),
        ("Acute Dyspnea / Stridor", "saans nahi aa rahi, breathless, dum ghut raha", "🔴 Oxygen & urgent ambulance"),
        ("Loss of Consciousness", "behosh, hosh nahi, unresponsive", "🔴 Recovery position + 108 dispatch"),
        ("Seizures / Convulsions", "daura, fits, mirgi, jhatke aana", "🔴 Protect airway + immediate transport"),
        ("Severe Hemorrhage", "bahut khoon, heavy bleeding", "🔴 Direct pressure + emergency care"),
        ("Pediatric Danger Signs", "unable to drink, lethargic child", "🔴 WHO IMNCI Red-alert referral")
    ]
    for condition, keywords, action in rules_table:
        p = tf.add_paragraph()
        p.text = f"• {condition}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = AMBER
        p = tf.add_paragraph()
        p.text = f"  Keywords: \"{keywords}\""
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p = tf.add_paragraph()
        p.text = f"  Action: {action}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ROSE
        p.space_after = Pt(4)

    # SLIDE 8
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Clinical Risk Stratification: Red / Yellow / Green Framework", "WHO IMNCI & NHM TRIAGE TIERS", 8)
    tiers = [
        ("🔴 RED: IMMEDIATE DANGER", ROSE, [
            ("Clinical Severity", "Life-threatening acute emergency requiring immediate intervention."),
            ("Trigger Conditions", "Chest pain, acute dyspnea, seizures, severe bleeding, altered consciousness."),
            ("System Protocol", "Bypasses LLM; locks screen to flashing red emergency card."),
            ("Action Required", "Instant 108 Ambulance dialer; immediate transfer to District Hospital / CHC.")
        ]),
        ("🟡 YELLOW: NEEDS ATTENTION", AMBER, [
            ("Clinical Severity", "Moderate clinical illness; self-care insufficient; medical evaluation needed."),
            ("Trigger Conditions", "Persistent high fever >3 days, productive cough, localized abdominal pain."),
            ("System Protocol", "Gemini 2.5 Flash analyzes vitals + symptoms + prescription context."),
            ("Action Required", "Visit nearest Primary Health Centre (PHC) within 24–48 hours; monitor vitals.")
        ]),
        ("🟢 GREEN: ROUTINE / STABLE", EMERALD, [
            ("Clinical Severity", "Mild, self-limiting condition suitable for home monitoring."),
            ("Trigger Conditions", "Mild cold, seasonal runny nose, minor fatigue with normal vital signs."),
            ("System Protocol", "Provides structured home remedies and clear warning signs for return."),
            ("Action Required", "Adequate hydration, rest, ORS if diarrhea; re-triage if fever escalates.")
        ])
    ]
    for idx, (t_title, t_color, t_items) in enumerate(tiers):
        t_left = Inches(0.8) + idx * Inches(3.95)
        create_card(s8, t_left, Inches(1.6), Inches(3.7), Inches(5.1), CARD_BG, t_color)
        tb = s8.shapes.add_textbox(t_left + Inches(0.2), Inches(1.8), Inches(3.3), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t_title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = t_color
        p.space_after = Pt(14)
        for i_head, i_desc in t_items:
            p = tf.add_paragraph()
            p.text = f"• {i_head}:"
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE
            p = tf.add_paragraph()
            p.text = f"  {i_desc}"
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(8)

    # SLIDE 9
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Vernacular Localization & Low-Literacy Inclusivity", "ACCESSIBILITY FOR 850M+ CITIZENS", 9)
    create_card(s9, Inches(0.8), Inches(1.6), Inches(5.66), Inches(5.1), CARD_BG, TEAL)
    tb = s9.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🌐 5-Language Instant Translation Engine"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(10)
    langs = [
        ("🇮🇳 Hindi (हिंदी) - DEFAULT", "Primary national language for Northern and Central states."),
        ("🇬🇧 English (English)", "Designed for clinicians, urban citizens, and hackathon evaluators."),
        ("🇮🇳 Bengali (বাংলা)", "Frontline language coverage for West Bengal, Tripura, and Assam."),
        ("🇮🇳 Telugu (తెలుగు)", "Community health coverage for Andhra Pradesh and Telangana."),
        ("🇮🇳 Marathi (मराठी)", "Field language coverage for Maharashtra PHCs and sub-centres.")
    ]
    for l_tit, l_sub in langs:
        p = tf.add_paragraph()
        p.text = f"✔ {l_tit}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"   {l_sub}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)
    p = tf.add_paragraph()
    p.text = "💡 Compact Header Switcher: Instant client-side state toggle with zero network latency and zero page reloads."
    p.font.size = Pt(11)
    p.font.color.rgb = EMERALD
    p.font.bold = True

    create_card(s9, Inches(6.86), Inches(1.6), Inches(5.66), Inches(5.1), CARD_BG, EMERALD)
    tb = s9.shapes.add_textbox(Inches(7.16), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🫀 Low-Literacy Inclusive Design"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(10)
    lit_pts = [
        ("Touch-Based Body Map", "Villagers unable to read or name organs simply tap where it hurts (Head, Chest, Stomach, Arms, Legs, Back)."),
        ("Audio-In, Audio-Out", "Illiterate patients don't have to decipher text — symptoms are spoken in, and guidance is spoken out."),
        ("Cognitive Color Coding", "Bold Red, Amber, and Green shield badges communicate urgency without requiring textual comprehension."),
        ("Acoustic Recording Feedback", "Pulsing concentric wave rings and equalizer bars give visual confidence that voice is being recorded.")
    ]
    for l_tit, l_sub in lit_pts:
        p = tf.add_paragraph()
        p.text = f"✔ {l_tit}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"   {l_sub}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    # SLIDE 10
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Multimodal Prescription OCR & Document Digitization", "VISION AI & DOCUMENT PROCESSING", 10)
    create_card(s10, Inches(0.8), Inches(1.6), Inches(5.66), Inches(5.1), CARD_BG, AMBER)
    tb = s10.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📸 The Rural Prescription Nightmare"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_after = Pt(10)
    ocr_challenges = [
        ("Illegible Doctor Cursive", "Rural patients carry crumpled paper slips with illegible handwritten dosages and drug names."),
        ("Lost Medical History", "Villagers lose paper slips between clinic visits, forcing doctors to guess past treatments."),
        ("Dangerous Drug Interactions", "ASHAs cannot decipher whether a patient is already taking Paracetamol, Antibiotics, or BP pills."),
        ("Multimodal Vision Extraction", "AarogyaVani allows ASHAs to snap a photo; Gemini Vision extracts medicines, dosages, and notes in <2s.")
    ]
    for tit, dsc in ocr_challenges:
        p = tf.add_paragraph()
        p.text = f"• {tit}:"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {dsc}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    create_card(s10, Inches(6.86), Inches(1.6), Inches(5.66), Inches(5.1), CARD_BG, TEAL)
    tb = s10.shapes.add_textbox(Inches(7.16), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚙️ Two-Tier Vision OCR & Verification"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(10)
    ocr_tech = [
        ("Tier 1: Gemini Multimodal Vision", "Processes raw images directly; deciphers Latin medical abbreviations (e.g., 'Tab PCM 500mg TDS', 'Amox 250 BD')."),
        ("Tier 2: Sarvam Doc-AI Fallback", "If Gemini is unavailable, automatically routes to Sarvam Doc-AI digitise API for Indic document parsing."),
        ("Animated Laser UI Scanner", "Green laser beam animation visually confirms image analysis on low-cost devices."),
        ("Human-in-the-Loop Correction", "Extracted text appears in an editable field. ASHA can correct misspelled patient names before triage.")
    ]
    for tit, dsc in ocr_tech:
        p = tf.add_paragraph()
        p.text = f"✔ {tit}:"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {dsc}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    # SLIDE 11
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Real-World Impact & National Health Mission Alignment", "IMPACT & PUBLIC HEALTH VALUE", 11)
    m_w = Inches(3.64)
    m_h = Inches(2.2)

    create_card(s11, Inches(0.8), Inches(1.6), m_w, m_h, CARD_BG, EMERALD)
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(1.75), m_w - Inches(0.4), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ <100ms"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p = tf.add_paragraph()
    p.text = "Emergency Response Time\nInstant client-side red-flag detection catches cardiac & respiratory crises immediately."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_LIGHT

    create_card(s11, Inches(4.84), Inches(1.6), m_w, m_h, CARD_BG, TEAL)
    tb = s11.shapes.add_textbox(Inches(5.04), Inches(1.75), m_w - Inches(0.4), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📉 70-80%"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p = tf.add_paragraph()
    p.text = "Reduction in Unnecessary Travel\nPrevents impoverished villagers from making 40km hospital trips for mild self-limiting colds."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_LIGHT

    create_card(s11, Inches(8.88), Inches(1.6), m_w, m_h, CARD_BG, AMBER)
    tb = s11.shapes.add_textbox(Inches(9.08), Inches(1.75), m_w - Inches(0.4), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📋 100%"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p = tf.add_paragraph()
    p.text = "Standardized Referrals\nEliminates vague verbal handoffs; PHC doctors receive structured clinical history & vitals."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_LIGHT

    create_card(s11, Inches(0.8), Inches(4.1), Inches(11.733), Inches(2.6), RGBColor(14, 24, 42), INDIGO)
    tb = s11.shapes.add_textbox(Inches(1.1), Inches(4.25), Inches(11.1), Inches(2.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🏛️ Alignment with Indian Government Digital Health Missions"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(8)
    gov_align = [
        ("National Health Mission (NHM)", "Adheres to ASHA doorstep screening protocols and WHO IMNCI clinical guidelines."),
        ("Ayushman Bharat Digital Mission (ABDM)", "Designed for seamless linkage to 14-digit ABHA IDs for longitudinal electronic health records."),
        ("Bhashini & Digital India", "Demonstrates sovereign AI capability using native Indian speech models and regional scripts.")
    ]
    for g_tit, g_desc in gov_align:
        p = tf.add_paragraph()
        p.text = f"✔ {g_tit}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = EMERALD
        p = tf.add_paragraph()
        p.text = f"   {g_desc}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT

    # SLIDE 12
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Roadmap, Scalability & Conclusion", "THE FUTURE OF RURAL HEALTHTECH", 12)

    create_card(s12, Inches(0.8), Inches(1.6), Inches(5.66), Inches(5.1), CARD_BG, TEAL)
    tb = s12.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🚀 Future Roadmap"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(12)
    road_pts = [
        ("1. Offline-First PWA (WASM)", "Embedding lightweight on-device speech models (Whisper TFLite/WASM) for zero-connectivity jungle & hill terrain."),
        ("2. ABHA ID Integration", "One-click generation and verification of 14-digit Ayushman Bharat Health Account IDs."),
        ("3. Automated PHC WhatsApp Alert", "Instantly dispatches digitized referral slip to the on-duty PHC Medical Officer ahead of patient arrival."),
        ("4. Epidemiological Heatmapping", "Anonymized aggregation of symptom clusters for district-level outbreak detection (e.g., Dengue, Malaria).")
    ]
    for r_tit, r_desc in road_pts:
        p = tf.add_paragraph()
        p.text = f"• {r_tit}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"  {r_desc}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    create_card(s12, Inches(6.86), Inches(1.6), Inches(5.66), Inches(5.1), RGBColor(14, 25, 45), EMERALD)
    tb = s12.shapes.add_textbox(Inches(7.16), Inches(1.8), Inches(5.06), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🌟 Key Takeaways for Judges"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(12)
    conc_pts = [
        ("Built for the Real India", "Engineered specifically for low-literacy rural patients and overburdened ASHA workers."),
        ("Safety is Uncompromised", "Deterministic gates ensure 100% zero-hallucination handling of acute medical emergencies."),
        ("Production-Ready & Tested", "Fully functional React 19 app, serverless endpoints, and passed E2E automated test suite."),
        ("Open-Source & Extensible", "Codebase published on GitHub with complete documentation and architecture whitepaper.")
    ]
    for c_tit, c_desc in conc_pts:
        p = tf.add_paragraph()
        p.text = f"✔ {c_tit}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p = tf.add_paragraph()
        p.text = f"   {c_desc}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "Repository: github.com/Prashantispacific/AROGYAVANI\nTeam LocalHost Boys • GDG HBTU • Bit N Build '26"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEAL

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AarogyaVani_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    build_presentation()
